from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_TEXT_CHARS = 50_000


def _truncate(text: str) -> str:
    return text[:MAX_TEXT_CHARS]


def extract_document_metadata(
    content: bytes,
    file_name: str,
    mime_type: str | None,
) -> dict[str, Any]:
    suffix = Path(file_name).suffix.lower()
    result: dict[str, Any] = {
        "file_name": file_name,
        "mime_type": mime_type,
        "extension": suffix,
        "byte_size": len(content),
        "extractor": "none",
        "text_preview": "",
    }

    try:
        if suffix == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            text = "
".join(page.extract_text() or "" for page in reader.pages)
            result.update(
                {
                    "extractor": "pypdf",
                    "page_count": len(reader.pages),
                    "text_preview": _truncate(text),
                }
            )
        elif suffix == ".docx":
            doc = DocxDocument(io.BytesIO(content))
            text = "
".join(paragraph.text for paragraph in doc.paragraphs)
            result.update(
                {
                    "extractor": "python-docx",
                    "paragraph_count": len(doc.paragraphs),
                    "text_preview": _truncate(text),
                }
            )
        elif suffix == ".xlsx":
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            parts: list[str] = []
            row_count = 0
            for sheet in wb.worksheets:
                parts.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    row_count += 1
                    parts.append(" | ".join("" if value is None else str(value) for value in row))
                    if sum(len(part) for part in parts) >= MAX_TEXT_CHARS:
                        break
                if sum(len(part) for part in parts) >= MAX_TEXT_CHARS:
                    break
            result.update(
                {
                    "extractor": "openpyxl",
                    "sheet_count": len(wb.sheetnames),
                    "row_count_scanned": row_count,
                    "text_preview": _truncate("
".join(parts)),
                }
            )
        elif suffix == ".csv":
            decoded = content.decode("utf-8-sig", errors="replace")
            reader = csv.reader(io.StringIO(decoded))
            rows: list[str] = []
            row_count = 0
            for row in reader:
                row_count += 1
                rows.append(" | ".join(row))
                if sum(len(part) for part in rows) >= MAX_TEXT_CHARS:
                    break
            result.update(
                {
                    "extractor": "csv",
                    "row_count_scanned": row_count,
                    "text_preview": _truncate("
".join(rows)),
                }
            )
    except Exception as exc:
        result["extraction_error"] = str(exc)

    return result
